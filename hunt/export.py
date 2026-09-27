"""Write tracker.xlsx and tracker.csv. Before writing, read back any status,
notes or date applied you typed into tracker.xlsx, so your edits are kept."""
import csv
import os
from datetime import datetime, timezone

from .config import STATUSES, TRACKER_CSV, TRACKER_XLSX, label_rank
from .db import loads, rows

COLUMNS = [
    ("id", "id"),
    ("date found", "date_found"),
    ("company", "company"),
    ("title", "title"),
    ("level", "level"),
    ("stretch?", "stretch"),
    ("country", "country"),
    ("location label", "location_label"),
    ("sponsorship", "sponsorship"),
    ("fit score", "fit_score"),
    ("why", "why"),
    ("gaps", "gaps"),
    ("apply link", "apply_url"),
    ("letter file", "letter_file"),
    ("cv file", "cv_file"),
    ("status", "status"),
    ("date applied", "date_applied"),
    ("notes", "notes"),
    ("source", "source"),
]
RESTRICTED_EXTRA = [("restriction (from the post)", "restriction")]
EDITABLE = {"status": "status", "notes": "notes", "date applied": "date_applied"}


def _cell(job, field):
    v = job.get(field)
    if field == "stretch":
        return "yes" if v else ""
    if field == "gaps":
        return ", ".join(loads(v))
    if field == "sponsorship":
        ev = loads(job.get("sponsorship_evidence"))
        return f"{v}: {'; '.join(ev)}" if ev else (v or "")
    if field == "fit_score":
        return round(v, 1) if v is not None else None
    return "" if v is None else v


def sorted_jobs(conn):
    jobs = rows(conn)
    jobs.sort(key=lambda j: (label_rank(j["location_label"]), -(j["fit_score"] or 0)))
    return jobs


def sync_from_xlsx(conn, path=TRACKER_XLSX, log=print):
    """Copy your edits in tracker.xlsx back into the database. An edit in
    the sheet wins unless you changed that job with `mark` after the sheet
    was last saved."""
    if not path.exists():
        return 0
    try:
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True, data_only=True)
    except Exception as e:  # noqa: BLE001
        log(f"  could not read {path.name} to keep your edits: {e}")
        return 0
    saved_at = datetime.fromtimestamp(os.path.getmtime(path), timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")
    changed = 0
    for ws in wb.worksheets:
        it = ws.iter_rows(values_only=True)
        header = next(it, None)
        if not header or "id" not in header:
            continue
        idx = {h: i for i, h in enumerate(header)}
        for r in it:
            if not r or r[idx["id"]] in (None, ""):
                continue
            job = conn.execute("SELECT * FROM jobs WHERE id = ?", (int(r[idx["id"]]),)).fetchone()
            if not job:
                continue
            if job["user_updated_at"] and job["user_updated_at"] > saved_at:
                continue
            updates = {}
            for col, field in EDITABLE.items():
                if col not in idx:
                    continue
                v = r[idx[col]]
                v = v.strftime("%Y-%m-%d") if hasattr(v, "strftime") else ("" if v is None else str(v).strip())
                if field == "status":
                    v = v.lower()
                    if v not in STATUSES:
                        continue
                if v != (job[field] or ""):
                    updates[field] = v or None
            if updates:
                sets = ", ".join(f"{k} = ?" for k in updates)
                conn.execute(f"UPDATE jobs SET {sets}, user_updated_at = ? WHERE id = ?",
                             [*updates.values(), saved_at, job["id"]])
                changed += 1
    conn.commit()
    if changed:
        log(f"  kept your edits from {path.name} for {changed} jobs")
    return changed


def export(conn, xlsx=TRACKER_XLSX, csv_path=TRACKER_CSV, log=print):
    jobs = sorted_jobs(conn)
    open_jobs = [j for j in jobs if j["location_label"] != "restricted"]
    restricted = [j for j in jobs if j["location_label"] == "restricted"]

    with open(csv_path, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        cols = COLUMNS + RESTRICTED_EXTRA
        w.writerow([c for c, _ in cols])
        for j in jobs:
            w.writerow([_cell(j, field) for _, field in cols])

    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    wb = Workbook()
    for i, (name, items, cols) in enumerate((("Open", open_jobs, COLUMNS),
                                             ("Restricted", restricted, COLUMNS + RESTRICTED_EXTRA))):
        ws = wb.active if i == 0 else wb.create_sheet()
        ws.title = name
        ws.append([c for c, _ in cols])
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="305496")
        for j in items:
            ws.append([_cell(j, field) for _, field in cols])
        link_col = [c for c, _ in cols].index("apply link") + 1
        for r in range(2, ws.max_row + 1):
            cell = ws.cell(r, link_col)
            if cell.value:
                cell.hyperlink = cell.value
                cell.font = Font(color="0563C1", underline="single")
        widths = {"id": 6, "date found": 11, "company": 18, "title": 38, "level": 11, "stretch?": 8, "country": 14,
                  "location label": 15, "sponsorship": 30, "fit score": 8, "why": 55, "gaps": 30, "apply link": 40,
                  "letter file": 30, "cv file": 30, "status": 10, "date applied": 12, "notes": 30, "source": 11,
                  "restriction (from the post)": 60}
        for k, (col, _) in enumerate(cols, start=1):
            ws.column_dimensions[ws.cell(1, k).column_letter].width = widths.get(col, 15)
        ws.freeze_panes = "E2"
        ws.auto_filter.ref = ws.dimensions
        status_col = ws.cell(1, [c for c, _ in cols].index("status") + 1).column_letter
        dv = DataValidation(type="list", formula1='"' + ",".join(STATUSES) + '"', allow_blank=False)
        ws.add_data_validation(dv)
        dv.add(f"{status_col}2:{status_col}{max(ws.max_row, 2) + 500}")
    try:
        wb.save(xlsx)
        out = xlsx
    except PermissionError:
        out = xlsx.with_name(f"tracker-{datetime.now():%Y%m%d-%H%M}.xlsx")
        wb.save(out)
        log(f"  {xlsx.name} is open in another program, saved to {out.name} instead")
    return out, len(open_jobs), len(restricted)
