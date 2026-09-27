from pathlib import Path

from hunt import db


def row(**kw):
    base = {"key": "k1", "url_key": "https://x/1", "apply_url": "https://x/1", "company": "Acme",
            "title": "ML Engineer", "location_label": "remote_open", "fit_score": 70.0, "why": "old why",
            "gaps": ["Spark"]}
    base.update(kw)
    return base


def test_rerun_never_overwrites_status_or_notes(tmp_path):
    conn = db.connect(tmp_path / "jobs.db")
    jid, new = db.upsert(conn, row(), today="2026-01-01")
    assert new
    db.set_status(conn, jid, "applied", note="sent via site")
    conn.execute("UPDATE jobs SET letter_file = 'output/letters/x.md' WHERE id = ?", (jid,))
    conn.commit()

    # A later run finds the same job with a new score.
    jid2, new2 = db.upsert(conn, row(fit_score=90.0, why="new why"), today="2026-02-01")
    conn.commit()
    assert jid2 == jid and not new2
    j = dict(conn.execute("SELECT * FROM jobs WHERE id = ?", (jid,)).fetchone())
    assert j["status"] == "applied"
    assert j["notes"] == "sent via site"
    assert j["letter_file"] == "output/letters/x.md"
    assert j["date_found"] == "2026-01-01"
    assert j["date_applied"] is not None
    assert j["fit_score"] == 90.0 and j["why"] == "new why"
    assert j["last_seen"] == "2026-02-01"


def test_same_url_different_key_updates_same_row(tmp_path):
    conn = db.connect(tmp_path / "jobs.db")
    jid, _ = db.upsert(conn, row())
    db.set_status(conn, jid, "interview")
    jid2, new = db.upsert(conn, row(key="k2", title="Machine Learning Engineer"))
    assert jid2 == jid and not new
    assert conn.execute("SELECT status FROM jobs WHERE id = ?", (jid,)).fetchone()[0] == "interview"


def test_set_letter_moves_new_to_drafted_only(tmp_path):
    conn = db.connect(tmp_path / "jobs.db")
    a, _ = db.upsert(conn, row())
    b, _ = db.upsert(conn, row(key="k2", url_key="https://x/2"))
    db.set_status(conn, b, "applied")
    db.set_letter(conn, a, "l1.md")
    db.set_letter(conn, b, "l2.md")
    assert conn.execute("SELECT status FROM jobs WHERE id = ?", (a,)).fetchone()[0] == "drafted"
    assert conn.execute("SELECT status FROM jobs WHERE id = ?", (b,)).fetchone()[0] == "applied"


def test_bad_status_is_refused(tmp_path):
    conn = db.connect(tmp_path / "jobs.db")
    jid, _ = db.upsert(conn, row())
    try:
        db.set_status(conn, jid, "maybe")
    except ValueError:
        return
    raise AssertionError("expected ValueError")


def test_edits_in_tracker_xlsx_are_kept(tmp_path):
    from openpyxl import load_workbook

    from hunt.export import export, sync_from_xlsx

    conn = db.connect(tmp_path / "jobs.db")
    jid, _ = db.upsert(conn, row())
    conn.commit()
    xlsx, csv_path = tmp_path / "tracker.xlsx", tmp_path / "tracker.csv"
    export(conn, xlsx=xlsx, csv_path=csv_path, log=lambda *_: None)

    # You type a status and a note straight into the sheet.
    wb = load_workbook(xlsx)
    ws = wb["Open"]
    header = [c.value for c in ws[1]]
    ws.cell(2, header.index("status") + 1).value = "skipped"
    ws.cell(2, header.index("notes") + 1).value = "not a fit"
    wb.save(xlsx)

    # The next run syncs first, then re-exports.
    sync_from_xlsx(conn, path=xlsx, log=lambda *_: None)
    db.upsert(conn, row(fit_score=10.0))
    export(conn, xlsx=xlsx, csv_path=csv_path, log=lambda *_: None)
    j = conn.execute("SELECT status, notes FROM jobs WHERE id = ?", (jid,)).fetchone()
    assert (j["status"], j["notes"]) == ("skipped", "not a fit")
    ws = load_workbook(xlsx)["Open"]
    assert ws.cell(2, header.index("status") + 1).value == "skipped"
    assert Path(csv_path).read_text().count("skipped") == 1


def test_drafted_status_survives_an_older_tracker_sheet(tmp_path):
    import os
    import time

    from hunt.export import export, sync_from_xlsx

    conn = db.connect(tmp_path / "jobs.db")
    jid, _ = db.upsert(conn, row())
    conn.commit()
    xlsx, csv_path = tmp_path / "tracker.xlsx", tmp_path / "tracker.csv"
    export(conn, xlsx=xlsx, csv_path=csv_path, log=lambda *_: None)
    old = time.time() - 120
    os.utime(xlsx, (old, old))  # the sheet was saved before the letter
    db.set_letter(conn, jid, "output/letters/x.md")
    sync_from_xlsx(conn, path=xlsx, log=lambda *_: None)
    assert conn.execute("SELECT status FROM jobs WHERE id = ?", (jid,)).fetchone()[0] == "drafted"


def test_an_old_sheet_does_not_undo_page_edits(tmp_path):
    import os
    import time

    from hunt.export import export, sync_from_xlsx

    conn = db.connect(tmp_path / "jobs.db")
    jid, _ = db.upsert(conn, row())
    conn.commit()
    db.set_letter(conn, jid, "output/letters/x.md")
    xlsx, csv_path = tmp_path / "tracker.xlsx", tmp_path / "tracker.csv"
    export(conn, xlsx=xlsx, csv_path=csv_path, log=lambda *_: None)  # sheet says drafted
    # You marked the job applied on the page before that export, and the
    # edit is imported afterwards with its own, older time.
    earlier = time.strftime("%Y-%m-%dT%H:%M:%SZ", time.gmtime(time.time() - 600))
    conn.execute("UPDATE jobs SET status = 'applied', user_updated_at = ? WHERE id = ?", (earlier, jid))
    conn.commit()
    sync_from_xlsx(conn, path=xlsx, log=lambda *_: None)
    assert conn.execute("SELECT status FROM jobs WHERE id = ?", (jid,)).fetchone()[0] == "applied"
    # Once you save the sheet yourself, its edits are read again.
    os.utime(xlsx, (time.time() + 5, time.time() + 5))
    sync_from_xlsx(conn, path=xlsx, log=lambda *_: None)
    assert conn.execute("SELECT status FROM jobs WHERE id = ?", (jid,)).fetchone()[0] == "drafted"
